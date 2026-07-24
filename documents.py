from __future__ import annotations

import hashlib
import io
import mimetypes
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument

try:
    import pymupdf
except ImportError:  # pragma: no cover
    import fitz as pymupdf  # type: ignore

from app.config import get_settings
from app.services.utils import safe_filename

settings = get_settings()

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class ExtractedPayload:
    filename: str
    mime_type: str
    raw_bytes: bytes
    pages: list[tuple[int | None, str]]



def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def infer_mime_type(filename: str, provided: str | None = None) -> str:
    if provided and provided not in {"application/octet-stream", "binary/octet-stream"}:
        return provided.split(";")[0].strip().lower()
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def extract_pdf(data: bytes) -> list[tuple[int | None, str]]:
    pdf = pymupdf.open(stream=data, filetype="pdf")
    try:
        return [(index + 1, page.get_text("text")) for index, page in enumerate(pdf)]
    finally:
        pdf.close()


def extract_docx(data: bytes) -> list[tuple[int | None, str]]:
    document = DocxDocument(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [(None, "\n".join(paragraphs))]


def extract_text(data: bytes) -> list[tuple[int | None, str]]:
    for encoding in ("utf-8", "latin-1"):
        try:
            return [(None, data.decode(encoding))]
        except UnicodeDecodeError:
            continue
    return [(None, data.decode("utf-8", errors="replace"))]


def extract_single(data: bytes, filename: str, mime_type: str | None = None) -> ExtractedPayload:
    filename = safe_filename(filename)
    suffix = Path(filename).suffix.lower()
    mime = infer_mime_type(filename, mime_type)

    if data.startswith(b"%PDF-") or suffix == ".pdf" or mime == "application/pdf":
        pages = extract_pdf(data)
        mime = "application/pdf"
        if suffix != ".pdf":
            filename = f"{filename}.pdf"
    elif suffix == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        pages = extract_docx(data)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif suffix in {".txt", ".md"} or mime.startswith("text/"):
        pages = extract_text(data)
    else:
        raise ValueError(f"Formato não suportado para extração: {filename} ({mime})")

    return ExtractedPayload(filename=filename, mime_type=mime, raw_bytes=data, pages=pages)


def extract_payloads(data: bytes, filename: str, mime_type: str | None = None) -> list[ExtractedPayload]:
    filename = safe_filename(filename)
    suffix = Path(filename).suffix.lower()
    mime = infer_mime_type(filename, mime_type)

    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if suffix == ".docx" or mime == docx_mime:
        return [extract_single(data, filename, mime)]
    if data.startswith(b"%PDF-"):
        return [extract_single(data, filename, "application/pdf")]

    zip_signature = zipfile.is_zipfile(io.BytesIO(data))
    if zip_signature:
        with zipfile.ZipFile(io.BytesIO(data)) as probe:
            if "word/document.xml" in probe.namelist():
                docx_name = filename if suffix == ".docx" else f"{filename}.docx"
                return [extract_single(data, docx_name, docx_mime)]

    if not zip_signature and (suffix == ".pdf" or mime == "application/pdf"):
        return [extract_single(data, filename, mime)]

    is_zip = (
        suffix == ".zip"
        or mime in {"application/zip", "application/x-zip-compressed"}
        or zip_signature
    )
    if not is_zip:
        return [extract_single(data, filename, mime)]

    payloads: list[ExtractedPayload] = []
    total_uncompressed = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = [member for member in archive.infolist() if not member.is_dir()]
        if len(members) > settings.max_zip_files:
            raise ValueError(f"ZIP contém mais de {settings.max_zip_files} arquivos")

        for member in members:
            member_path = Path(member.filename)
            if member_path.is_absolute() or ".." in member_path.parts:
                continue
            suffix = member_path.suffix.lower()
            if suffix not in SUPPORTED_EXTENSIONS:
                continue
            if member.file_size > settings.max_document_bytes:
                raise ValueError(
                    f"Arquivo {member_path.name} no ZIP excede o limite de "
                    f"{settings.max_document_mb} MB"
                )
            total_uncompressed += member.file_size
            if total_uncompressed > settings.max_zip_uncompressed_bytes:
                raise ValueError("ZIP excede o limite de dados descompactados")
            member_data = archive.read(member)
            payloads.append(extract_single(member_data, member_path.name))

    if not payloads:
        raise ValueError("O ZIP não contém PDF, DOCX, TXT ou Markdown suportado")
    return payloads


def save_payload(edital_id: str, document_id: str, payload: ExtractedPayload) -> Path:
    target_dir = settings.data_dir / edital_id
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{document_id}_{safe_filename(payload.filename)}"
    target.write_bytes(payload.raw_bytes)
    return target
